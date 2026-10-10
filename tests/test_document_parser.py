# -*- coding: utf-8 -*-
"""文档格式试卷解析测试（PDF / docx）。

背景：家长上传的试卷不只是照片，也可能是电子版 PDF / Word。
- 文本型 PDF/docx 应直接提取文本（零 LLM 成本、100% 准确），不走 vision；
- 扫描型 PDF（无文本层）应渲染成图片回退 vision OCR；
- 老格式 .doc 无法解析，应提示「另存为 .docx」。
"""

import os

import pytest


@pytest.fixture
def env(test_db_path, demo_mode, monkeypatch, tmp_path):
    """隔离 uploads 目录（不污染项目 uploads/）。"""
    from pipeline import stages
    upload_dir = str(tmp_path / "uploads")
    monkeypatch.setattr(stages, "UPLOAD_DIR", upload_dir)
    return upload_dir


def _make_docx(path, paragraphs, table=None):
    import docx
    d = docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for i, row in enumerate(table):
            for j, val in enumerate(row):
                t.cell(i, j).text = val
    d.save(path)
    return path


def _make_pdf(path, text=None):
    """text 非空 → 文本型 PDF；None → 空白页（模拟扫描件，无文本层）。"""
    import fitz
    doc = fitz.open()
    page = doc.new_page()
    if text:
        page.insert_text((72, 100), text)
    doc.save(path)
    doc.close()
    return path


def test_docx_extract_text_and_table(tmp_path):
    """docx 段落 + 表格内容都能提取（试卷常用表格排版）。"""
    from document_parser import extract_text_from_docx
    p = _make_docx(str(tmp_path / "paper.docx"),
                   ["31. I volunteered as ___ guide."],
                   table=[["题号", "答案"], ["31", "B"]])
    text = extract_text_from_docx(p)
    assert "volunteered" in text
    assert "B" in text  # 表格内容


def test_text_pdf_extracted_directly(tmp_path):
    """文本型 PDF 直接提取，无需 vision。"""
    from document_parser import extract_text_from_pdf
    p = _make_pdf(str(tmp_path / "t.pdf"),
                  "31. I volunteered as ___ guide in Chengdu Museum." * 2)
    text, ok = extract_text_from_pdf(p)
    assert ok is True
    assert "volunteered" in text


def test_scanned_pdf_falls_back_to_images(tmp_path):
    """扫描型 PDF（无文本层）→ 渲染图片回退 vision。"""
    from document_parser import extract_text_from_pdf, pdf_to_images
    p = _make_pdf(str(tmp_path / "s.pdf"))
    text, ok = extract_text_from_pdf(p)
    assert ok is False
    imgs = pdf_to_images(p, str(tmp_path / "out"))
    assert len(imgs) == 1
    assert os.path.exists(imgs[0])


def test_is_document_and_legacy():
    from document_parser import is_document, is_legacy_doc
    assert is_document("a.pdf") and is_document("b.docx")
    assert not is_document("c.jpg")
    assert is_legacy_doc("d.doc")


def _save_file(student_id, src, name, file_type, upload_dir):
    import db
    d = os.path.join(upload_dir, str(student_id), file_type)
    os.makedirs(d, exist_ok=True)
    with open(src, "rb") as f:
        data = f.read()
    with open(os.path.join(d, name), "wb") as f:
        f.write(data)
    return db.add_file(student_id=student_id, uploader_role="parent",
                       file_type=file_type, filename=name,
                       original_filename=name, week_start=db.get_week_start(),
                       file_size=len(data), mime_type="application/octet-stream")


def test_ocr_docx_skips_vision(env, sample_student, test_db_path,
                               monkeypatch, tmp_path):
    """docx 走文档解析，不调用 vision（省成本且 100% 准确）。"""
    from pipeline import stages
    import skills_bridge

    called = []

    def _spy(paths, task_id=None):
        called.extend(paths)
        return [{"ok": True, "text": "VISION", "confidence": 0.9}
                for _ in paths]

    monkeypatch.setattr(skills_bridge, "run_ocr_parallel", _spy)

    src = _make_docx(str(tmp_path / "p.docx"),
                     ["31. I volunteered as ___ guide."])
    ids = [_save_file(sample_student, src, "p.docx", "test_paper", env)]

    text, _conf = stages.ocr_to_text(ids, sample_student, task_id=1,
                                     db_path=test_db_path)
    assert called == [], "docx 不应走 vision OCR"
    assert "volunteered" in text
    assert "--- 第 1 页 ---" in text


def test_ocr_scanned_pdf_uses_vision(env, sample_student, test_db_path,
                                     monkeypatch, tmp_path):
    """扫描型 PDF 渲染成图片后走 vision。"""
    from pipeline import stages
    import skills_bridge

    called = []

    def _spy(paths, task_id=None):
        called.extend(paths)
        return [{"ok": True, "text": "VISION_TEXT", "confidence": 0.9}
                for _ in paths]

    monkeypatch.setattr(skills_bridge, "run_ocr_parallel", _spy)

    src = _make_pdf(str(tmp_path / "s.pdf"))
    ids = [_save_file(sample_student, src, "s.pdf", "test_paper", env)]

    text, _conf = stages.ocr_to_text(ids, sample_student, task_id=1,
                                     db_path=test_db_path)
    assert len(called) == 1, "扫描型 PDF 应渲染成图片走 vision"
    assert called[0].endswith(".png")
    assert "VISION_TEXT" in text


def test_validate_upload_doc_hint():
    """.doc 给出「另存为 .docx」提示；.docx / .pdf / 图片放行。"""
    from web.shared import _validate_upload

    class _F:
        def __init__(self, name, size=10):
            self.filename = name
            self.content_length = size

    err = _validate_upload(_F("a.doc"))
    assert err and "docx" in err
    assert _validate_upload(_F("b.docx")) is None
    assert _validate_upload(_F("c.pdf")) is None
    assert _validate_upload(_F("d.jpg")) is None
    assert _validate_upload(_F("e.exe")) is not None
